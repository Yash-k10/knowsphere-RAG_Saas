import os
from pathlib import Path
import pymupdf as fitz  # PyMuPDF
import docx

class ExtractionError(Exception):
    pass

class TextExtractor:
    @staticmethod
    def extract(file_path: str | Path, file_type: str) -> list[dict]:
        """
        Extracts structured text from supported file types (PDF, DOCX, TXT).
        Returns a list of page/segment dictionaries:
        [{"page_number": 1, "text": "Page content..."}, ...]
        """
        path = Path(file_path)
        if not path.exists():
            raise ExtractionError(f"File not found: {file_path}")

        ext = path.suffix.lower()
        if ext == ".pdf":
            return TextExtractor._extract_pdf(path)
        elif ext == ".docx":
            return TextExtractor._extract_docx(path)
        elif ext == ".txt":
            return TextExtractor._extract_txt(path)
        else:
            raise ExtractionError(f"Unsupported file format: {ext}. Only PDF, DOCX, and TXT are supported.")

    @staticmethod
    def _extract_pdf(path: Path) -> list[dict]:
        pages = []
        try:
            doc = fitz.open(str(path))
            for page_idx in range(len(doc)):
                page = doc.load_page(page_idx)
                text = page.get_text("text").strip()
                if text:
                    pages.append({
                        "page_number": page_idx + 1,
                        "text": text
                    })
            doc.close()
            if not pages:
                raise ExtractionError("PDF file contains no extractable text (it might be scanned images).")
            return pages
        except Exception as e:
            raise ExtractionError(f"Error parsing PDF file: {str(e)}")

    @staticmethod
    def _extract_docx(path: Path) -> list[dict]:
        try:
            doc = docx.Document(str(path))
            full_text = []
            for para in doc.paragraphs:
                cleaned = para.text.strip()
                if cleaned:
                    full_text.append(cleaned)
            
            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        full_text.append(row_text)

            combined = "\n\n".join(full_text)
            if not combined.strip():
                raise ExtractionError("DOCX file contains no readable text.")

            return [{
                "page_number": None,
                "text": combined
            }]
        except Exception as e:
            raise ExtractionError(f"Error parsing DOCX file: {str(e)}")

    @staticmethod
    def _extract_txt(path: Path) -> list[dict]:
        try:
            for encoding in ["utf-8", "latin-1", "cp1252"]:
                try:
                    with open(path, "r", encoding=encoding) as f:
                        text = f.read().strip()
                        if text:
                            return [{"page_number": 1, "text": text}]
                except UnicodeDecodeError:
                    continue
            raise ExtractionError("TXT file is empty or cannot be decoded with standard encodings.")
        except Exception as e:
            raise ExtractionError(f"Error reading TXT file: {str(e)}")
